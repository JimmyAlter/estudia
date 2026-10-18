"""Arranque de Estudia: servidor local, navegador y sincronización periódica.

Todo lo que tiene efectos va dentro de `main()`: la verificación con SymPy usa
un subproceso que vuelve a importar este módulo (spawn en Windows).
"""
import logging
import multiprocessing
import os
import signal
import socket
import sys
import threading
import time
import webbrowser
from functools import partial
from pathlib import Path

import uvicorn

from estudia.config import data_dir
from estudia.db import connect as db_connect
from estudia.server.app import HOST, create_app, periodic_sync
from estudia.launcher import acquire_lock, headless, opens_browser, running_instance
from estudia.server.handler import ensure_installed, is_estudia
from estudia.reminders import run_reminders
from estudia.streams import redirect_missing_streams
from estudia.tray import run_tray, shutdown, tray_available

log = logging.getLogger(__name__)

STARTUP_WAIT_S = 15
INSTANCE_WAIT_S = 10  # espera a que la otra instancia, que ya tiene el lock, termine de arrancar
GRACEFUL_SHUTDOWN_S = 3  # uvicorn corta los pedidos y tareas de fondo que sigan después de esto
SERVER_JOIN_S = 5


def free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind((HOST, 0))
        return s.getsockname()[1]


def uvicorn_config(app, port: int) -> uvicorn.Config:
    # Sin colores: sin consola no hay terminal que los entienda (ver estudia.streams).
    # Con un límite al apagar: una tarea de fondo larga (primera sincronización, Gemini)
    # no puede dejar a "Salir" esperando para siempre.
    return uvicorn.Config(app, host=HOST, port=port, log_level="warning", use_colors=False,
                          timeout_graceful_shutdown=GRACEFUL_SHUTDOWN_S)


def _open_when_ready(server: uvicorn.Server, url: str) -> None:
    deadline = time.monotonic() + STARTUP_WAIT_S
    while not server.started and time.monotonic() < deadline:
        time.sleep(0.1)
    if server.started:
        webbrowser.open(url)


def _reuse(url: str) -> None:
    if opens_browser():
        webbrowser.open(url + "/")


def _wait_for_instance(port_file, lock_path) -> tuple[str | None, object | None]:
    """Espera a la instancia que tiene el lock: (su URL, None), o (None, el lock) si murió."""
    deadline = time.monotonic() + INSTANCE_WAIT_S
    while True:
        url = running_instance(port_file, is_estudia)
        if url:
            return url, None
        # La otra pudo cerrarse o caerse antes de escribir el puerto.
        lock = acquire_lock(lock_path)
        if lock is not None or time.monotonic() >= deadline:
            return None, lock
        time.sleep(0.2)


BUSY_TEXT = (
    "Estudia ya se está abriendo o no responde. Cerrala desde la bandeja "
    "o el Administrador de tareas y volvé a abrirla."
)


def _message_box(text: str) -> None:
    import ctypes

    ctypes.windll.user32.MessageBoxW(None, text, "Estudia", 0x40)


CRASH_TEXT = "Estudia no pudo arrancar. Volvé a abrirla; si sigue pasando, avisanos en GitHub."


def _fail(text: str) -> None:
    """Avisa y sale con error."""
    print(text, file=sys.stderr)
    # Sin consola (la app de Windows no la tiene) el aviso se ve en un cuadro nativo.
    if sys.platform == "win32" and not headless():
        _message_box(text)
    sys.exit(1)


def _report_busy() -> None:
    """El lock lo tiene otra instancia que no contesta."""
    _fail(BUSY_TEXT)


CHECK_WORKER_FLAG = "--check-worker"


# Plantilla lineal de prueba para el modo de variantes: f(x) = 2x + 3, ¿f(4)? Está escrita acá
# (no se importa de los tests) para que `--check-worker` la pruebe en el .exe. Un test comprueba
# que coincide con `tests/variant_fixtures.py`.
_CHECK_VARIANTS_PAYLOAD = {
    "params": [{"name": "a", "low": 2, "high": 9}, {"name": "b", "low": 1, "high": 9},
               {"name": "c", "low": 2, "high": 9}],
    "original": [2, 3, 4],
    "answer_expr": "a*c + b",
    "distractor_exprs": ["a*(c + b)", "a + b", "a*c"],
    "explanation": "f({c}) = {a}·{c} + {b} = {respuesta}.",
    "common_error": "Olvidarse de sumar {b} a {a}·{c}.",
    "check_expr": "11",
    "original_options": ["11", "14", "5", "8"],
    "original_explanation": "f(4) = 2·4 + 3 = 11.",
    "original_common_error": "Olvidarse de sumar 3 a 2·4.",
    "correct_index": 0,
    "fractions": False,
    "seed": "check-worker",
    "limit": 4,
}


def _worker_works() -> bool:
    """El worker de SymPy arranca (forkserver o spawn, también congelado), corrige bien y calcula
    variantes."""
    from estudia.generation import verify

    try:
        return (verify.check_math_tristate("sqrt(2)", "√2") is True
                and verify.grade_typed("sqrt(2)", "1.41") is True
                and verify.grade_typed("sqrt(2)", "1.4") is False
                and bool(verify.compute_variant_values(_CHECK_VARIANTS_PAYLOAD)))
    except Exception:
        return False


def check_worker() -> int:
    """`Estudia.exe --check-worker`: 0 si el worker de SymPy anda, 1 si no (lo usa la prueba de CI)."""
    ok = _worker_works()
    print("worker de SymPy: " + ("bien" if ok else "no anda"))
    return 0 if ok else 1


CHECK_DOCUMENTS_FLAG = "--check-documents"


def check_documents() -> int:
    """`Estudia.exe --check-documents`: 0 si se leen PDF, diapositivas y fotos; 1 si no.

    Lo usa la prueba de CI del .exe: prueba la DLL de pdfium, el proceso hijo congelado, las
    plantillas de python-pptx, Pillow y python-multipart. Imprime solo el paso que falló.
    """
    import io
    import tempfile

    from estudia.material import documents, files

    step = "preparar"
    try:
        with tempfile.TemporaryDirectory(prefix="estudia-check-") as tmp:
            step = "crear un PDF con pdfium"
            import pypdfium2

            pdf = pypdfium2.PdfDocument.new()
            for _ in range(2):
                pdf.new_page(200, 200)
            pdf_path = Path(tmp) / "prueba.pdf"
            pdf.save(str(pdf_path))
            pdf.close()

            step = "leer el PDF"
            info = documents.inspect_document(pdf_path, files.PDF)
            if info.pages != 2:
                raise ValueError("páginas")

            step = "cortar una página del PDF"
            if not documents.pdf_pages_bytes(pdf_path, 2, 2).startswith(b"%PDF"):
                raise ValueError("pdf")

            step = "crear diapositivas con python-pptx"
            from pptx import Presentation

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Prueba de Estudia"
            pptx_path = Path(tmp) / "prueba.pptx"
            prs.save(str(pptx_path))

            step = "leer las diapositivas"
            if "Prueba de Estudia" not in documents.slides_text(pptx_path):
                raise ValueError("pptx")

            step = "limpiar una foto"
            from PIL import Image

            buf = io.BytesIO()
            Image.new("RGB", (8, 8), (200, 100, 50)).save(buf, "JPEG")
            if not files.reencode_image(buf.getvalue(), files.JPEG):
                raise ValueError("jpeg")

            step = "cargar python-multipart"
            import python_multipart  # noqa: F401 - FastAPI lo necesita para recibir archivos
    except Exception:
        print("Lectura de documentos: no anda (" + step + ")")
        return 1
    print("Lectura de documentos: bien")
    return 0


CHECK_VIDEO_TOOLS_FLAG = "--check-video-tools"
JS_CHECK_TIMEOUT_S = 30


def _run_js(name: str, path: str) -> str:
    """Corre un script mínimo con el runtime de JS, como lo hace yt-dlp: sin ventana."""
    import subprocess

    script = "console.log('ok')"
    args = [path, "eval", script] if name == "deno" else [path, "-e", script]
    flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
    proc = subprocess.run(args, capture_output=True, text=True, timeout=JS_CHECK_TIMEOUT_S,
                          stdin=subprocess.DEVNULL, creationflags=flags)
    return proc.stdout.strip() if proc.returncode == 0 else ""


def check_video_tools() -> int:
    """`Estudia.exe --check-video-tools`: 0 si yt-dlp y su runtime de JS andan; 1 si no.

    Sin red: arma el YoutubeDL con las opciones de `material.video`, corre el runtime de JS
    y le pide a yt-dlp que lo acepte y que cargue con él los scripts del desafío de YouTube.
    En el ejecutable congelado exige el runtime incluido. Fuera de él usa el que usaría
    `material.video` (el deno de la dependencia o node); si no hay ninguno, avisa y sale con 0.
    Imprime solo el paso que falló.
    """
    import shutil

    from estudia import config
    from estudia.material import video

    step = "cargar yt-dlp"
    try:
        import yt_dlp

        options = video.ydl_options("check.%(ext)s")
        with yt_dlp.YoutubeDL(options) as ydl:
            step = "encontrar el runtime de JavaScript"
            if getattr(sys, "frozen", False):
                bundled = config.bundled_js_runtime_path()
                if bundled is None or not bundled.is_file():
                    raise FileNotFoundError("runtime")
                name, path = config.BUNDLED_JS_RUNTIME, str(bundled)
                if options["js_runtimes"] != {name: {"path": path}}:
                    raise ValueError("runtime")
            else:
                found = [(n, (c or {}).get("path") or shutil.which(n))
                         for n, c in options["js_runtimes"].items()]
                found = [(n, p) for n, p in found if p and Path(p).is_file()]
                if not found:
                    print("Herramientas de video: sin runtime empaquetado (modo desarrollo)")
                    return 0
                name, path = found[0]

            step = "correr el runtime de JavaScript"
            if _run_js(name, path) != "ok":
                raise RuntimeError("js")

            step = "que yt-dlp acepte el runtime de JavaScript"
            runtime = ydl._js_runtimes.get(name)
            if runtime is None or runtime.info is None or not runtime.info.supported:
                raise RuntimeError("runtime")

            step = "resolver con yt-dlp el desafío de YouTube"
            from yt_dlp.extractor.youtube.jsc._director import initialize_jsc_director

            director = initialize_jsc_director(ydl.get_info_extractor("Youtube"))
            provider = next(p for p in director.providers.values()
                            if getattr(p, "JS_RUNTIME_NAME", None) == name)
            if not provider.is_available():
                raise RuntimeError("provider")
            # Los scripts del desafío (yt_dlp_ejs) corridos con el runtime y los permisos que
            # usa yt-dlp: sin el reproductor de YouTube, alcanza con que definan `jsc`.
            code = (f"{provider._lib_script.code}\nObject.assign(globalThis, lib);\n"
                    f"{provider._core_script.code}\nconsole.log(typeof jsc);\n")
            if provider._run_js_runtime(code).strip() != "function":
                raise RuntimeError("jsc")
    except Exception:
        print("Herramientas de video: no anda (" + step + ")")
        return 1
    print("Herramientas de video: bien")
    return 0


def warn_if_worker_broken() -> None:
    """Al arrancar: un solo aviso en el log si el worker de SymPy no anda (sin detalles)."""
    if not _worker_works():
        log.warning("No arrancó el verificador de matemática (SymPy): las respuestas escritas "
                    "y la verificación de preguntas nuevas no van a andar.")


def main() -> None:
    # El login vuelve por `estudia://` o `moodlemobile://`: en Linux se registra solo.
    ensure_installed()
    port_file = data_dir() / "port"
    url = running_instance(port_file, is_estudia)
    if url:
        _reuse(url)
        return
    # Un doble clic rápido: la primera instancia tiene el lock pero aún no escribió el puerto.
    lock_path = data_dir() / "instance.lock"
    lock = acquire_lock(lock_path)
    if lock is None:
        url, lock = _wait_for_instance(port_file, lock_path)
        if url:
            _reuse(url)
            return
        if lock is None:
            _report_busy()
    port = free_port()
    port_file.write_text(str(port), "utf-8")
    # Una sola base para el servidor y los recordatorios.
    db_path = data_dir() / "estudia.db"
    app = create_app(db_path=db_path, port=port)
    server = uvicorn.Server(uvicorn_config(app, port))

    # uvicorn vuelve a lanzar SIGTERM al terminar; sin esto el proceso muere
    # antes del `finally` y queda el archivo del puerto.
    signal.signal(signal.SIGTERM, lambda *_: sys.exit(0))
    stop = threading.Event()
    threading.Thread(target=periodic_sync, args=(app.state.jobs, stop), daemon=True).start()
    threading.Thread(target=warn_if_worker_broken, name="worker-check", daemon=True).start()
    if not headless():
        threading.Thread(target=run_reminders, args=(stop, partial(db_connect, db_path)), daemon=True).start()
    page_url = f"http://{HOST}:{port}/"
    use_tray = not headless()
    if use_tray and not tray_available():
        log.warning("No se pudo cargar la bandeja del sistema; Estudia sigue sin ícono.")
        use_tray = False
    if opens_browser():
        threading.Thread(target=_open_when_ready, args=(server, page_url), daemon=True).start()
    server_thread = None
    exit_code = 0
    try:
        if use_tray:
            # pystray exige el hilo principal en Windows, así que el servidor va en otro.
            server_thread = threading.Thread(target=server.run, name="uvicorn")
            server_thread.start()
            try:
                run_tray(page_url, lambda: shutdown(server), lambda: not server_thread.is_alive())
            except Exception:
                # Sin ícono se puede seguir usando desde el navegador: no se corta el servidor.
                log.exception("La bandeja dejó de andar; Estudia sigue sin ícono.")
                while server_thread.is_alive():
                    server_thread.join(0.5)
            if not server.should_exit:  # el hilo murió solo: error de arranque o caída
                _fail(CRASH_TEXT)
        else:
            server.run()
    except SystemExit as exc:
        if not use_tray:
            raise
        exit_code = exc.code if isinstance(exc.code, int) else int(exc.code is not None)
    finally:
        shutdown(server)
        if server_thread is not None:
            server_thread.join(SERVER_JOIN_S)
        stop.set()
        # Un archivo viejo haría que el handler mande el token a otro proceso.
        port_file.unlink(missing_ok=True)
        lock.close()
    if use_tray:
        _hard_exit(exit_code)


def _hard_exit(code: int) -> None:
    """Termina ya, sin esperar hilos que sigan en un trabajo largo.

    Si no, el proceso queda vivo sin ícono ni servidor. SQLite aguanta un corte así.
    """
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.flush()
        except Exception:
            pass
    for handler in logging.getLogger().handlers:
        handler.flush()
    os._exit(code)


if __name__ == "__main__":
    multiprocessing.freeze_support()
    # Estudia.exe no tiene consola: lo que se imprima (uvicorn, errores) va al log.
    redirect_missing_streams(data_dir() / "estudia.log")
    if CHECK_WORKER_FLAG in sys.argv[1:]:
        sys.exit(check_worker())
    if CHECK_DOCUMENTS_FLAG in sys.argv[1:]:
        sys.exit(check_documents())
    if CHECK_VIDEO_TOOLS_FLAG in sys.argv[1:]:
        sys.exit(check_video_tools())
    main()
