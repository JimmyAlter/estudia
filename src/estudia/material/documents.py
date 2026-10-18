"""Leer PDF y diapositivas con topes: páginas, texto, marcadores y partes.

Todo lo que abre el archivo corre en `sandbox.run_limited` (otro proceso, con tiempo y memoria
limitados). Las funciones que corren en el hijo reciben los topes como argumentos: el hijo
importa este módulo de cero y no ve lo que cambie el proceso principal.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

from estudia.material.files import (
    PDF,
    PPTX,
    TOO_LARGE_TEXT,
    UNREADABLE_TEXT,
    TooLarge,
    UnreadableFile,
    check_pptx_zip,
)
from estudia.material.sandbox import SandboxBusy, SandboxFailed, SandboxTimeout, run_limited

PDF_MAX_PAGES = 1000
PAST_EXAM_MAX_PAGES = 30
SLIDES_MAX = 500
SPLIT_OVER_PAGES = 60
CHUNK_PAGES = 40
MIN_PART_PAGES = 5
MIN_TAIL_PAGES = 10
TEXT_SAMPLE_PAGES = 30
TEXT_PAGE_MIN_CHARS = 50
TOC_MAX_ITEMS = 300
PPTX_MAX_CHARS = 300_000
PARSE_TIMEOUT_S = 60

_UNREADABLE = UNREADABLE_TEXT
_TOO_LARGE = TOO_LARGE_TEXT


class DocumentTimeout(UnreadableFile):
    """Leer el archivo se pasó del tiempo."""


class DocumentBusy(Exception):
    """Todos los lugares para leer estaban ocupados: hay que probar más tarde (no es el archivo)."""


class TooManyPages(TooLarge):
    """El documento tiene más páginas (o diapositivas) que el tope."""


@dataclass
class DocInfo:
    kind: str
    pages: int  # páginas, o diapositivas en un .pptx
    has_text: bool
    outline: list[tuple[str, int]]  # marcadores de primer nivel: (título, página desde 1)


@dataclass
class Part:
    title: str | None
    page_from: int  # desde 1, inclusivo
    page_to: int  # inclusivo


# --- en el hijo ----------------------------------------------------------------------------


def _sample_indexes(count: int) -> list[int]:
    """Hasta TEXT_SAMPLE_PAGES índices repartidos parejo por todo el documento."""
    if count <= TEXT_SAMPLE_PAGES:
        return list(range(count))
    return sorted({i * count // TEXT_SAMPLE_PAGES for i in range(TEXT_SAMPLE_PAGES)})


def _has_text(lengths: list[int]) -> bool:
    """False si la mitad o más de la muestra son páginas escaneadas (casi sin texto)."""
    if not lengths:
        return False
    short = sum(1 for n in lengths if n < TEXT_PAGE_MIN_CHARS)
    return short * 2 < len(lengths)


def _inspect_pdf(path: Path, max_pages: int) -> DocInfo:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        pages = len(pdf)
        if pages > max_pages:
            raise TooManyPages(_TOO_LARGE)
        lengths = []
        for index in _sample_indexes(pages):
            page = pdf[index]
            textpage = page.get_textpage()
            lengths.append(len(textpage.get_text_range().strip()))
            textpage.close()
            page.close()
        outline = []
        for visited, item in enumerate(pdf.get_toc(max_depth=1)):
            if visited >= TOC_MAX_ITEMS:  # cuenta también los que no tienen destino
                break
            dest = item.get_dest()
            index = dest.get_index() if dest is not None else None
            if index is None:
                continue
            outline.append(((item.get_title() or "").strip(), index + 1))
        return DocInfo(kind=PDF, pages=pages, has_text=_has_text(lengths), outline=outline)
    finally:
        pdf.close()


def _shape_texts(shapes) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_shape_texts(shape.shapes))
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                texts.append(text)
        elif shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    texts.append(" | ".join(cells))
    return texts


def _slide_parts(slide) -> tuple[str, str]:
    body = "\n".join(_shape_texts(slide.shapes))
    notes = ""
    if slide.has_notes_slide:
        frame = slide.notes_slide.notes_text_frame
        notes = frame.text.strip() if frame is not None else ""
    return body, notes


def _inspect_pptx(path: Path, max_slides: int) -> DocInfo:
    from pptx import Presentation

    slides = Presentation(str(path)).slides
    count = len(slides)
    if count > max_slides:
        raise TooManyPages(_TOO_LARGE)
    lengths = []
    for index in _sample_indexes(count):
        body, notes = _slide_parts(slides[index])
        lengths.append(len(body) + len(notes))
    return DocInfo(kind=PPTX, pages=count, has_text=_has_text(lengths), outline=[])


def _slides_text(
    path: Path, page_from: int, page_to: int | None, max_chars: int, max_slides: int
) -> str:
    from pptx import Presentation

    slides = Presentation(str(path)).slides
    if len(slides) > max_slides:
        raise TooManyPages(_TOO_LARGE)
    last = len(slides) if page_to is None else min(page_to, len(slides))
    blocks: list[str] = []
    size = 0
    for number in range(page_from, last + 1):
        body, notes = _slide_parts(slides[number - 1])
        # Numeradas desde 1 en lo que se manda, como el PDF de un tramo: la tubería las corre.
        block = f"Diapositiva {number - page_from + 1}:\n{body}"
        if notes:
            block += f"\nNotas: {notes}"
        blocks.append(block)
        size += len(block) + 2
        if size >= max_chars:
            break
    return "\n\n".join(blocks)[:max_chars]


def _pdf_pages(path: Path, page_from: int, page_to: int, max_pages: int) -> bytes:
    import pypdfium2 as pdfium

    src = pdfium.PdfDocument(str(path))
    try:
        if len(src) > max_pages:
            raise TooManyPages(_TOO_LARGE)
        if page_to > len(src):
            raise ValueError("rango fuera del documento")
        out = pdfium.PdfDocument.new()
        try:
            out.import_pages(src, list(range(page_from - 1, page_to)))
            buf = io.BytesIO()
            out.save(buf)
            return buf.getvalue()
        finally:
            out.close()
    finally:
        src.close()


# Claves de metadatos de un PDF (diccionario /Info y XMP del catálogo). Un PDF reescrito por
# `strip_pdf_metadata` no puede tener ninguna: si queda alguna, no se guarda.
_METADATA_KEYS = (b"/Author", b"/Creator", b"/Producer", b"/Title", b"/Subject", b"/Keywords",
                  b"/CreationDate", b"/ModDate", b"/Trapped", b"/Metadata")
_TRAILER_INFO = re.compile(rb"/Info\s+(\d+)\s+0\s+R")


def _blank_info(pdf: bytes) -> bytes:
    """Vacía el diccionario /Info que pdfium escribe al guardar (Creator, Producer y la fecha):
    lo reemplaza por `<<  >>` con espacios del mismo largo, así el xref sigue valiendo."""
    refs = _TRAILER_INFO.findall(pdf)
    if not refs:
        return pdf
    obj = re.compile(rb"(?m)^%s 0 obj\s*<<(.*?)>>\s*endobj" % refs[-1], re.S)
    match = obj.search(pdf)
    if match is None:
        raise ValueError("no se encontró el diccionario /Info")
    start, end = match.span(1)
    return pdf[:start] + b" " * (end - start) + pdf[end:]


# Un parcial en PDF se guarda rasterizado: cada página es una foto JPEG nueva. Es lo más seguro
# contra los metadatos: no sobrevive nada del archivo original salvo los píxeles (ni /Info, ni
# XMP del documento o de una imagen, ni el EXIF/GPS de una foto incrustada, ni anotaciones con
# datos). Se pierde la capa de texto: la IA lee el parcial como imagen, igual que una foto.
# 150 ppp con calidad 85 se lee bien (un A4 queda en 1240 × 1754) y 30 páginas entran holgadas
# en el tope de 20 MB. El lado más largo se acota para que una página gigante no agote la memoria.
RASTER_DPI = 150
RASTER_MAX_SIDE_PX = 2500
RASTER_JPEG_QUALITY = 85
# Marcas de metadatos de una imagen JPEG: el JPEG nuevo (de Pillow, sin `exif=`) no tiene ninguna.
_IMAGE_METADATA_MARKS = (b"Exif\x00", b"xmpmeta", b"ns.adobe.com/xap")


def strip_pdf_metadata(data: bytes, max_pages: int) -> bytes:
    """El PDF rasterizado con pdfium: cada página se dibuja (con sus anotaciones) a RASTER_DPI y
    vuelve como una foto JPEG nueva, del mismo tamaño de página. Corre en el sandbox. ValueError
    si quedó alguna marca de metadatos: un parcial nunca se guarda con ellos."""
    import pypdfium2 as pdfium

    src = pdfium.PdfDocument(data)
    try:
        if len(src) > max_pages:
            raise TooManyPages(_TOO_LARGE)
        out = pdfium.PdfDocument.new()
        try:
            for index in range(len(src)):
                page = src[index]
                try:
                    width, height = page.get_size()
                    if width <= 0 or height <= 0:
                        raise ValueError("página sin tamaño")
                    scale = min(RASTER_DPI / 72, RASTER_MAX_SIDE_PX / max(width, height))
                    image = page.render(scale=scale, may_draw_forms=True).to_pil().convert("RGB")
                finally:
                    page.close()
                jpeg = io.BytesIO()
                image.save(jpeg, "JPEG", quality=RASTER_JPEG_QUALITY, optimize=True)
                jpeg.seek(0)
                new_page = out.new_page(width, height)
                photo = pdfium.PdfImage.new(out)
                photo.load_jpeg(jpeg, pages=[new_page], inline=True)
                photo.set_matrix(pdfium.PdfMatrix().scale(width, height))
                new_page.insert_obj(photo)
                new_page.gen_content()
                new_page.close()
            buf = io.BytesIO()
            out.save(buf)
        finally:
            out.close()
    finally:
        src.close()
    clean = _blank_info(buf.getvalue())
    if any(mark in clean for mark in _METADATA_KEYS + _IMAGE_METADATA_MARKS):
        raise ValueError("el PDF reescrito todavía tiene metadatos")
    return clean


# --- en el proceso principal ---------------------------------------------------------------


def _run(target, *args, timeout_s: float):
    try:
        return run_limited(target, *args, timeout_s=timeout_s)
    except SandboxBusy as exc:
        raise DocumentBusy("ocupado") from exc
    except SandboxTimeout as exc:
        raise DocumentTimeout(_UNREADABLE) from exc
    except SandboxFailed as exc:
        if str(exc) == "TooManyPages":
            raise TooManyPages(_TOO_LARGE) from exc
        raise UnreadableFile(_UNREADABLE) from exc


def inspect_document(
    path: Path,
    kind: str,
    *,
    max_pages: int = PDF_MAX_PAGES,
    timeout_s: float | None = None,
) -> DocInfo:
    """Páginas, si tiene texto y marcadores. `timeout_s` None: PARSE_TIMEOUT_S."""
    timeout = PARSE_TIMEOUT_S if timeout_s is None else timeout_s
    if kind == PDF:
        return _run(_inspect_pdf, Path(path), max_pages, timeout_s=timeout)
    if kind == PPTX:
        check_pptx_zip(path)
        return _run(_inspect_pptx, Path(path), min(max_pages, SLIDES_MAX), timeout_s=timeout)
    raise ValueError(f"tipo de documento no soportado: {kind}")


def _chunks(title: str | None, page_from: int, page_to: int) -> list[Part]:
    """Tramos de CHUNK_PAGES; un último tramo de menos de MIN_TAIL_PAGES se une al anterior."""
    ranges = []
    start = page_from
    while start <= page_to:
        end = min(start + CHUNK_PAGES - 1, page_to)
        if ranges and end - start + 1 < MIN_TAIL_PAGES:
            ranges[-1] = (ranges[-1][0], end)
        else:
            ranges.append((start, end))
        start = end + 1
    if title is None:
        return [Part(None, a, b) for a, b in ranges]
    if len(ranges) == 1:
        return [Part(title, *ranges[0])]
    base = title.strip() or "Parte"
    return [Part(f"{base} ({i})", a, b) for i, (a, b) in enumerate(ranges, start=1)]


def plan_parts(info: DocInfo) -> list[Part]:
    """Partes de un documento largo: por marcadores o en tramos. [] si no hace falta partirlo."""
    pages = info.pages
    if pages <= SPLIT_OVER_PAGES:
        return []
    marks: list[tuple[str, int]] = []
    seen_pages: set[int] = set()
    for title, page in sorted(info.outline, key=lambda item: item[1]):
        if 1 <= page <= pages and page not in seen_pages:
            seen_pages.add(page)
            marks.append((title, page))
    if len(marks) < 2:
        return _chunks(None, 1, pages)

    # [título, desde, hasta]: la primera arranca en la 1 y la última termina al final.
    parts = []
    for i, (title, page) in enumerate(marks):
        start = 1 if i == 0 else page
        end = marks[i + 1][1] - 1 if i + 1 < len(marks) else pages
        parts.append([title, start, end])

    merged: list[list] = []
    pending = None  # parte corta que espera unirse a la siguiente
    for title, start, end in parts:
        if pending is not None:
            start = pending[1]
            pending = None
        if end - start + 1 < MIN_PART_PAGES:
            pending = [title, start, end]
            continue
        merged.append([title, start, end])
    if pending is not None:
        if merged:
            merged[-1][2] = pending[2]
        else:
            merged.append(pending)

    result: list[Part] = []
    for title, start, end in merged:
        if end - start + 1 > SPLIT_OVER_PAGES:
            result.extend(_chunks(title, start, end))
        else:
            result.append(Part(title, start, end))
    return result


def pdf_pages_bytes(
    path: Path,
    page_from: int,
    page_to: int,
    *,
    timeout_s: float | None = None,
) -> bytes:
    """Un PDF nuevo con las páginas `page_from..page_to` (desde 1, inclusivo)."""
    if page_from < 1 or page_to < page_from:
        raise ValueError("rango de páginas inválido")
    if page_to - page_from + 1 > SPLIT_OVER_PAGES:
        raise ValueError("rango de páginas demasiado largo")
    timeout = PARSE_TIMEOUT_S if timeout_s is None else timeout_s
    return _run(_pdf_pages, Path(path), page_from, page_to, PDF_MAX_PAGES, timeout_s=timeout)


def slides_text(
    path: Path,
    page_from: int = 1,
    page_to: int | None = None,
    *,
    timeout_s: float | None = None,
) -> str:
    """Texto de las diapositivas, con sus notas, cortado en PPTX_MAX_CHARS.

    Cada una va rotulada desde 1 dentro de lo pedido: en un tramo `page_from..page_to`, la
    `page_from` es la "Diapositiva 1".
    """
    check_pptx_zip(path)
    timeout = PARSE_TIMEOUT_S if timeout_s is None else timeout_s
    return _run(
        _slides_text,
        Path(path),
        max(page_from, 1),
        page_to,
        PPTX_MAX_CHARS,
        SLIDES_MAX,
        timeout_s=timeout,
    )
