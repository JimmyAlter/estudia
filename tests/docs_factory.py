"""Generador de archivos sintéticos diminutos para los tests (PDF, pptx, fotos, bombas)."""

import io
import struct
import zlib

from PIL import Image


def _esc(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _pdf_text(text: str) -> str:
    """Cadena de PDF: literal si entra en latin-1; si no, UTF-16BE con BOM en hexa (como los PDF reales)."""
    try:
        text.encode("latin-1")
    except UnicodeEncodeError:
        return "<FEFF" + text.encode("utf-16-be", "surrogatepass").hex().upper() + ">"
    return f"({_esc(text)})"


def _latin(text: str) -> bytes:
    return text.encode("latin-1", "replace")


def make_pdf(
    pages: list[str],
    *,
    outline: list[tuple[str, int]] = (),
    inflate_mb: int = 0,
    encrypt: bool = False,
    info: dict[str, str] | None = None,
    xmp: bytes | None = None,
) -> bytes:
    """PDF 1.4 escrito a mano: Helvetica, un `Tj` por página y xref correcto.

    `outline` son marcadores de primer nivel (título, número de página desde 1). `encrypt` suma
    un diccionario `/Encrypt` estándar (RC4, R2) cuyas claves no aceptan la contraseña vacía:
    hace falta una contraseña de usuario para abrirlo. `info` suma un diccionario `/Info`
    (p. ej. {"Author": "Ana"}) y `xmp`, un stream `/Metadata` en el catálogo.
    """
    n = len(pages)
    objs: dict[int, bytes] = {}
    font_id = 3
    first_page = 4
    first_content = first_page + n
    outline_root = first_content + n if outline else None
    kids = " ".join(f"{first_page + i} 0 R" for i in range(n))
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode()
    objs[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i, text in enumerate(pages):
        objs[first_page + i] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {first_content + i} 0 R >>"
        ).encode()
        body = _latin(f"BT /F1 12 Tf 72 720 Td ({_esc(text)}) Tj ET")
        if i == 0 and inflate_mb:
            # El stream inflado son espacios, y el texto de la página va al final.
            comp = zlib.compressobj(9)
            data = bytearray()
            chunk = b" " * (1024 * 1024)
            for _ in range(inflate_mb):
                data += comp.compress(chunk)
            data += comp.compress(body) + comp.flush()
            data = bytes(data)
            objs[first_content + i] = (
                b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(data)
                + data
                + b"\nendstream"
            )
        else:
            objs[first_content + i] = (
                b"<< /Length %d >>\nstream\n" % len(body) + body + b"\nendstream"
            )
    root_extra = ""
    if outline:
        items = list(outline)
        first_item = outline_root + 1
        objs[outline_root] = (
            f"<< /Type /Outlines /First {first_item} 0 R "
            f"/Last {first_item + len(items) - 1} 0 R /Count {len(items)} >>"
        ).encode()
        for j, (title, page) in enumerate(items):
            parts = [f"/Title {_pdf_text(title)}", f"/Parent {outline_root} 0 R"]
            if j > 0:
                parts.append(f"/Prev {first_item + j - 1} 0 R")
            if j < len(items) - 1:
                parts.append(f"/Next {first_item + j + 1} 0 R")
            parts.append(f"/Dest [{first_page + page - 1} 0 R /Fit]")
            objs[first_item + j] = _latin("<< " + " ".join(parts) + " >>")
        root_extra = f" /Outlines {outline_root} 0 R"
    trailer_extra = b""
    if xmp is not None:
        xmp_id = max(objs) + 1
        objs[xmp_id] = (b"<< /Type /Metadata /Subtype /XML /Length %d >>\nstream\n" % len(xmp)
                        + xmp + b"\nendstream")
        root_extra += f" /Metadata {xmp_id} 0 R"
    if info:
        info_id = max(objs) + 1
        objs[info_id] = _latin("<< " + " ".join(f"/{k} {_pdf_text(v)}" for k, v in info.items()) + " >>")
        trailer_extra += b" /Info %d 0 R" % info_id
    objs[1] = f"<< /Type /Catalog /Pages 2 0 R{root_extra} >>".encode()
    if encrypt:
        encrypt_id = max(objs) + 1
        objs[encrypt_id] = (
            b"<< /Filter /Standard /V 1 /R 2 /Length 40 /P -4 /O <" + b"ab" * 32 + b"> /U <"
            + b"cd" * 32 + b"> >>"
        )
        trailer_extra += b" /Encrypt %d 0 R /ID [<%s> <%s>]" % (encrypt_id, b"01" * 16, b"01" * 16)

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n"
    xref = len(out)
    size = max(objs) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size
    for num in range(1, size):
        out += b"%010d 00000 n \n" % offsets[num]
    out += b"trailer\n<< /Size %d /Root 1 0 R%s >>\nstartxref\n%d\n%%%%EOF\n" % (
        size, trailer_extra, xref)
    return bytes(out)


def blank_pdf(n: int) -> bytes:
    """PDF con `n` páginas vacías."""
    return make_pdf([""] * n)


def make_pptx(slides: list[str], notes: list[str] | None = None) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[5]  # solo título
    for i, text in enumerate(slides):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = text
        if notes and i < len(notes) and notes[i]:
            slide.notes_slide.notes_text_frame.text = notes[i]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _photo_exif(gps: bool) -> Image.Exif:
    exif = Image.Exif()
    exif[0x010F] = "CamaraSecreta"  # Make
    if gps:
        exif[0x8825] = {1: "S", 2: (34.0, 36.0, 0.0), 3: "W", 4: (58.0, 22.0, 0.0)}
    return exif


def make_jpeg(size=(64, 48), gps=True, comment: bytes | None = b"ComentarioSecreto") -> bytes:
    """JPEG con EXIF (marca de cámara y, si `gps`, coordenadas) y un comentario COM."""
    img = Image.new("RGB", size, (200, 80, 40))
    buf = io.BytesIO()
    extra = {"comment": comment} if comment else {}
    img.save(buf, "JPEG", exif=_photo_exif(gps), **extra)
    return buf.getvalue()


def make_mpo(size=(64, 48)) -> bytes:
    """Foto MPO como las de algunos celulares: dos JPEG (con EXIF y GPS) unidos por un MPF."""
    first = Image.new("RGB", size, (200, 80, 40))
    second = Image.new("RGB", size, (10, 200, 10))
    buf = io.BytesIO()
    first.save(buf, "MPO", save_all=True, append_images=[second], exif=_photo_exif(True))
    return buf.getvalue()


def make_png(size=(64, 48), text_chunk="secreto") -> bytes:
    from PIL import PngImagePlugin

    img = Image.new("RGB", size, (40, 120, 200))
    info = PngImagePlugin.PngInfo()
    if text_chunk:
        info.add_text("nota", text_chunk)
    buf = io.BytesIO()
    img.save(buf, "PNG", pnginfo=info)
    return buf.getvalue()


def _chunk(tag: bytes, data: bytes) -> bytes:
    crc = zlib.crc32(tag + data) & 0xFFFFFFFF
    return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", crc)


def png_claiming(width: int, height: int) -> bytes:
    """PNG cuyo IHDR declara `width` x `height`, con un IDAT mínimo."""
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + _chunk(b"IHDR", ihdr)
        + _chunk(b"IDAT", zlib.compress(b"\x00"))
        + _chunk(b"IEND", b"")
    )


def zip_bomb_pptx() -> bytes:
    """Esqueleto pptx más una entrada de 400 MB declarados, de ceros comprimidos."""
    import zipfile

    skeleton = zipfile.ZipFile(io.BytesIO(make_pptx(["x"])))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in skeleton.namelist():
            zf.writestr(name, skeleton.read(name))
        with zf.open("ppt/media/bomba.bin", "w", force_zip64=True) as entry:
            chunk = b"\0" * (1024 * 1024)
            for _ in range(400):
                entry.write(chunk)
    return buf.getvalue()


def make_image_pdf(jpeg: bytes, *, metadata: bytes | None = None, size=(64, 48)) -> bytes:
    """PDF de una página con una foto JPEG incrustada tal cual (DCTDecode, con su EXIF), como un
    escáner o un celular. `metadata` suma un stream XMP `/Metadata` al objeto de la imagen."""
    width, height = size
    objs: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /XObject << /Im1 5 0 R >> >> /Contents 4 0 R >>"),
    }
    body = b"q 400 0 0 300 100 400 cm /Im1 Do Q"
    objs[4] = b"<< /Length %d >>\nstream\n" % len(body) + body + b"\nendstream"
    meta_ref = b""
    if metadata is not None:
        objs[6] = (b"<< /Type /Metadata /Subtype /XML /Length %d >>\nstream\n" % len(metadata)
                   + metadata + b"\nendstream")
        meta_ref = b" /Metadata 6 0 R"
    objs[5] = (b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceRGB"
               b" /BitsPerComponent 8 /Filter /DCTDecode /Length %d%s >>\nstream\n"
               % (width, height, len(jpeg), meta_ref) + jpeg + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n"
    xref = len(out)
    size_ = max(objs) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size_
    for num in range(1, size_):
        out += b"%010d 00000 n \n" % offsets[num]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size_, xref)
    return bytes(out)
